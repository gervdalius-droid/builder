import openpyxl, sys, warnings
warnings.filterwarnings('ignore')
A=openpyxl.load_workbook('/home/claude/rep/ref.xlsm',keep_vba=True)
B=openpyxl.load_workbook(sys.argv[1] if len(sys.argv)>1 else 'test/rep/rebuilt.xlsm',keep_vba=True)
def val(c):
    v=c.value
    if hasattr(v,'text'): v=v.text
    return v
tot=0
for name,cols,first in [('PARTLIST',range(1,19),1),('LABEL',range(1,15),1),('EDGEBAND',range(2,15),1),('STOCK_CORE',range(2,21),1),('HARDWARE_LIST',range(2,9),1)]:
    a,b=A[name],B[name]
    last=max(r for r in range(1,a.max_row+1) if any(a.cell(r,c).value not in (None,'') for c in cols))
    lastb=max(r for r in range(1,b.max_row+1) if any(b.cell(r,c).value not in (None,'') for c in cols))
    diffs=[]; hdiff=[]
    for r in range(first,max(last,lastb)+1):
        for c in cols:
            va,vb=val(a.cell(r,c)),val(b.cell(r,c))
            if isinstance(va,str) and va.startswith('='): va=va.replace('$','')
            if isinstance(vb,str) and vb.startswith('='): vb=vb.replace('$','')
            if (va if va!='' else None)!=(vb if vb!='' else None): diffs.append((a.cell(r,c).coordinate,va,vb))
            if a.cell(r,c).style_id!=b.cell(r,c).style_id and r<=last: hdiff.append((a.cell(r,c).coordinate,'style',a.cell(r,c).style_id,b.cell(r,c).style_id))
        if a.row_dimensions[r].height!=b.row_dimensions[r].height and r<=last: hdiff.append((r,'height',a.row_dimensions[r].height,b.row_dimensions[r].height))
    ma=set(str(m) for m in a.merged_cells.ranges if m.min_row<=last); mb=set(str(m) for m in b.merged_cells.ranges)
    print(f'{name:14s} rows {last}/{lastb}  value diffs {len(diffs)}  style/height diffs {len(hdiff)}  merges only-ref {len(ma-mb)} only-new {len(mb-ma)}')
    for d in diffs[:8]: print('    ',d)
    for d in hdiff[:5]: print('    ',d)
    for m in sorted(ma-mb)[:5]: print('     merge missing',m)
    for m in sorted(mb-ma)[:5]: print('     merge extra',m)
    tot+=len(diffs)+len(hdiff)
print('TOTAL', tot)
